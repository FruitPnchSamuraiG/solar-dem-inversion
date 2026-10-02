"""Shared helpers for building decks on the NYU template (slides/DEM.pptx).

Everything written here passes the OOXML schema, which Google Slides enforces:
integer coordinates, table borders before fills, positive chart axis ids.
"""
import copy
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDES = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(SLIDES)
IMG = os.path.join(SLIDES, "img")
P13 = os.path.join(REPO, "results/plots/13_bright_diagnostic_20260928")

PURPLE = RGBColor(0x57, 0x06, 0x8C)
DEEP = RGBColor(0x33, 0x06, 0x62)
DARK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
CARD = RGBColor(0xF4, 0xF1, 0xF8)
LAV = RGBColor(0xE3, 0xDF, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xB9, 0xB4, 0xC2)
LIGHT = RGBColor(0xEC, 0xEB, 0xE8)

# Common geometry (16:9, 10 x 5.625 in); the NYU logo sits below 4.85 in.
L = Inches(0.4)
TOP = Inches(1.0)
FULL_W = Inches(9.2)
BOTTOM = Inches(4.85)


class Deck:
    """The NYU template with its placeholder content slide removed."""

    def __init__(self, template=os.path.join(SLIDES, "DEM.pptx")):
        self.prs = Presentation(template)
        ids = self.prs.slides._sldIdLst
        for sid in list(ids)[1:]:
            self.prs.part.drop_rel(sid.get(qn("r:id")))
            ids.remove(sid)
        self.blank = next(l for l in self.prs.slide_layouts if l.name == "BLANK")
        self.big = next(l for l in self.prs.slide_layouts if l.name == "BIG_NUMBER")
        self.sldnum = _sldnum(self.blank)

    @property
    def title_slide(self):
        return self.prs.slides[0]

    def slide(self, title, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        for sp in self.sldnum:
            s.shapes._spTree.append(copy.deepcopy(sp))
        text(s, Inches(0.29), Inches(0.24), Inches(9.3), Inches(0.62), [title],
             size=20, color=PURPLE, bold=True)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def divider(self, part, name):
        """Section divider on the template's BIG_NUMBER layout: "Part N" over the section name."""
        s = self.prs.slides.add_slide(self.big)
        for ph in list(s.placeholders):
            if ph.placeholder_format.type != PP_PLACEHOLDER.TITLE:
                ph._element.getparent().remove(ph._element)
        for sp in _sldnum(self.big):
            s.shapes._spTree.append(copy.deepcopy(sp))
        title = s.shapes.title
        title.left, title.top, title.width, title.height = 311700, 846450, 8520600, 1671000  # as placed in Google Slides
        p = title.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        for i, words in enumerate((f"Part {part}", name)):
            if i:
                p.add_line_break()
            r = p.add_run()
            r.text = words
            r.font.size = Pt(41)
        return s

    def save(self, path):
        self.prs.save(path)
        return path


def _sldnum(layout):
    """The layout's slide-number placeholder, which python-pptx does not clone onto new slides."""
    return [sp for sp in layout.shapes._spTree.iterchildren(qn("p:sp"))
            if sp.find(".//" + qn("p:ph")) is not None
            and sp.find(".//" + qn("p:ph")).get("type") == "sldNum"]


def text(slide, x, y, w, h, paras, size=13, color=DARK, bold=False, bullets=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6, italic=False, font=None):
    """paras: list of str, or of lists of (str, {bold, color, size, italic, font}) runs."""
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for txt, opt in runs:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", italic)
            f.color.rgb = opt.get("color", color)
            if opt.get("font", font):
                f.name = opt.get("font", font)
            if opt.get("sub"):
                r._r.get_or_add_rPr().set("baseline", "-25000")
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.2))))
            pPr.set("indent", str(-Emu(Inches(0.2))))
            clr = etree.SubElement(etree.SubElement(pPr, qn("a:buClr")), qn("a:srgbClr"))
            clr.set("val", str(color))
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    return tb


def picture(slide, path, x, y, w, h, align="center"):
    """Fit the image inside the box, preserving aspect ratio."""
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    px = x + (w - pw) // 2 if align == "center" else x
    py = y + (h - ph) // 2
    return slide.shapes.add_picture(path, int(px), int(py), pw, ph)


def card(slide, x, y, w, h, fill=CARD):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(x), int(y), int(w), int(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def box(slide, x, y, w, h, head, sub="", fill=CARD, head_color=PURPLE, sub_color=DARK,
        head_size=12, sub_size=9.5):
    """A rounded box with a bold head line and an optional smaller line."""
    b = card(slide, x, y, w, h, fill=fill)
    tf = b.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = head
    r.font.size = Pt(head_size); r.font.bold = True; r.font.color.rgb = head_color
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(sub_size); r2.font.color.rgb = sub_color
    return b


def table(slide, x, y, w, rows, col_w=None, size=10.5, highlight=(), row_h=Inches(0.3),
          text_cols=1, bold_first_col=False):
    nr, nc = len(rows), len(rows[0])
    shape = slide.shapes.add_table(nr, nc, int(x), int(y), int(w), int(row_h * nr))
    t = shape.table
    tblPr = t._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        tblPr.remove(style)
    if col_w:
        total = sum(col_w)
        for j, cw in enumerate(col_w):
            t.columns[j].width = int(w * cw / total)
    for i, row in enumerate(rows):
        t.rows[i].height = int(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = PURPLE if i == 0 else (LAV if i in highlight else WHITE)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j < text_cols else PP_ALIGN.RIGHT
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            r.font.bold = (i == 0) or (j == 0 and (bold_first_col or i in highlight))
            r.font.color.rgb = WHITE if i == 0 else DARK
            # Borders must precede the fill in a:tcPr; Google Slides rejects the reverse.
            tcPr = c._tc.get_or_add_tcPr()
            for k, edge in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                ln = etree.Element(qn(edge))
                ln.set("w", "6350" if edge == "a:lnB" else "0")
                if edge == "a:lnB":
                    etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", "D9D3E3")
                else:
                    etree.SubElement(ln, qn("a:noFill"))
                tcPr.insert(k, ln)
    return shape


def line(slide, x1, y1, x2, y2, color=PURPLE, width=1.5, arrow=True, dashed=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = MSO_LINE.DASH
    if arrow:
        etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")
    return c


def caption(slide, x, y, w, words):
    text(slide, x, y, w, Inches(0.3), [words], size=9.5, color=MUTED, italic=True, space_after=0)


def fix_chart_axis_ids(chart):
    """python-pptx can write negative axis ids; the schema needs unsignedInt."""
    ids = {}
    for el in chart._chartSpace.iter(qn("c:axId"), qn("c:crossAx")):
        v = el.get("val")
        if int(v) < 0:
            el.set("val", ids.setdefault(v, str(500000001 + len(ids))))
